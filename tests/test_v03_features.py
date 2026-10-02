"""Tests for v0.3.0: markets sweep, search-engine discovery parsing,
Excel export, challenge-resume plumbing."""

from lead_scout.markets import STATE_ABBRS_ALL, sweep_queries, us_markets
from lead_scout.models import Lead
from lead_scout.providers.search_engine_discovery import (
    _dedupe_companies,
    find_company_urls_via_search,
)
from lead_scout.storage import write_xlsx


def test_us_markets_cover_50_states_plus_dc():
    markets = us_markets()
    assert len(markets) == 51
    assert "Texas" in markets and "District of Columbia" in markets
    assert STATE_ABBRS_ALL["Texas"] == "TX"
    assert STATE_ABBRS_ALL["District of Columbia"] == "DC"


def test_sweep_queries_template():
    queries = sweep_queries("Total Wireless retailer")
    assert len(queries) == 51
    assert queries[0] == "Total Wireless retailer in Alabama"
    assert "Total Wireless retailer in Texas" in queries


def test_dedupe_companies_extracts_and_skips_chrome():
    html = ('<a href="https://www.linkedin.com/company/acme-wireless/">x</a>'
            '<a href="https://linkedin.com/company/acme-wireless/">dup</a>'
            '<a href="https://www.linkedin.com/company/feed/">chrome</a>')
    urls = _dedupe_companies(html, limit=10)
    assert urls == ["https://www.linkedin.com/company/acme-wireless/"]


def test_search_returns_gracefully_without_network(monkeypatch):
    # engines unreachable -> (empty list, status string), never raises
    def boom(*_a, **_k):
        raise OSError("offline")
    import lead_scout.providers.search_engine_discovery as mod
    monkeypatch.setattr(mod, "_read", boom)
    urls, status = find_company_urls_via_search("total wireless", limit=5)
    assert urls == []
    assert isinstance(status, str) and status


def test_write_xlsx_creates_two_sheets(tmp_path):
    import openpyxl
    target = tmp_path / "leads.xlsx"
    lead = Lead(business_name="Acme Wireless", owner_name="Owner Name",
                phone="713 555 0100", status="ok")
    blocked = Lead(business_name="Blocked Dealer", status="http_403_not_bypassed")
    write_xlsx(target, [lead, blocked])
    book = openpyxl.load_workbook(target)
    assert book.sheetnames == ["Leads", "Manual Review"]
    assert book["Leads"].max_row == 3          # header + 2 rows
    assert book["Leads"].cell(row=2, column=1).value == "Acme Wireless"
    assert book["Manual Review"].max_row == 2  # header + 1 blocked row
