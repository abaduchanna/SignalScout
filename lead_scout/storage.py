from __future__ import annotations

import csv
from pathlib import Path

from .extract import normalized_domain
from .models import Lead

FIELDS = (
    "business_name", "owner_name", "phone", "email", "address", "website",
    "linkedin_company_url", "source_type", "source_urls", "evidence", "status",
)


def read_seeds(path: str | Path) -> list[Lead]:
    with Path(path).open(newline="", encoding="utf-8-sig") as handle:
        return [Lead(**{key: value for key, value in row.items() if key in Lead.__dataclass_fields__})
                for row in csv.DictReader(handle)]


def dedupe(leads: list[Lead]) -> list[Lead]:
    merged: dict[str, Lead] = {}
    for lead in leads:
        domain = normalized_domain(lead.website) if lead.website else ""
        key = domain or f"{lead.business_name.lower()}|{lead.address.lower()}"
        if not key.strip("|"):
            key = f"unknown-{len(merged)}"
        merged.setdefault(key, Lead()).merge(lead)
    return list(merged.values())


def write_csv(path: str | Path, leads: list[Lead]) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    with target.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS)
        writer.writeheader()
        for lead in leads:
            writer.writerow(lead.as_csv_row())

def write_xlsx(path: str | Path, leads: list[Lead]) -> None:
    """Excel workbook (user request: results in .xlsx, not just CSV).

    Sheet "Leads" = every record; sheet "Manual Review" = non-ok records
    (blocked, challenged, incomplete). Brand-styled header, frozen top
    row, sensible column widths."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    def _fill_sheet(sheet, rows: list[Lead]) -> None:
        header_fill = PatternFill("solid", fgColor="0D0C14")
        header_font = Font(color="6EE7EF", bold=True, name="Consolas",
                           size=9)
        for col, field in enumerate(FIELDS, start=1):
            cell = sheet.cell(row=1, column=col, value=field)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center")
        for row_index, lead in enumerate(rows, start=2):
            data = lead.as_csv_row()
            for col, field in enumerate(FIELDS, start=1):
                sheet.cell(row=row_index, column=col,
                           value=data.get(field, ""))
        widths = (30, 22, 16, 28, 42, 34, 38, 16, 46, 46, 24)
        for col, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(col)].width = width
        sheet.freeze_panes = "A2"

    book = Workbook()
    _fill_sheet(book.active, list(leads))          # sheet: Leads
    book.active.title = "Leads"
    review = book.create_sheet("Manual Review")
    _fill_sheet(review, [lead for lead in leads if lead.status != "ok"])
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    book.save(target)
